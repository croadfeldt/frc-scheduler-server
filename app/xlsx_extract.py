"""Parse FMS-style schedule XLSX files back into match lists.

Mirrors the layout produced by `_buildXlsxWorksheet` in static/index.html
so we can round-trip the app's own exports. The format is canonical
FMS:

    Row 1: Title (e.g. "Qualification Match Schedule")
    Row 2: blank
    Row 3: Subtitle (event line)
    Row 4: Matches Per Team | (blank) | <int>
    Row 5: Time | Description | Blue 1 | Blue 2 | Blue 3 | Red 1 | Red 2 | Red 3
    Row 6+: data rows

The Description column carries the match number ("Qualification 12",
"Practice 3"). Surrogate notation is NOT preserved by the FMS export
format — we mark every surrogate flag false on import. If a user
needs surrogates they can edit the preview before committing.

Multi-sheet workbooks are common: the app exports a Qualification
sheet and (if present) a Practice sheet. We parse both — qual matches
become the primary match list; practice matches are returned
separately so the UI can decide what to do with them. Today both
flow into the import pipeline as straight matches; the practice
distinction is preserved in the returned dict.
"""
from __future__ import annotations

import logging
import re
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

log = logging.getLogger(__name__)

# Title-row text → "kind" tag. Matches against case-insensitive
# substrings so "Qualification Match Schedule", "Quals", "Qualifying"
# all map to "qual".
SHEET_KIND_HEURISTICS: list[tuple[str, str]] = [
    ("practice",      "practice"),
    ("qualification", "qual"),
    ("qual",          "qual"),
]

# Matches "Qualification 12", "Practice 3", "Q12", "P3", or just "12"
# in the Description column. Number is the only thing we actually need.
_MATCH_NUM_RE = re.compile(r"(\d+)")


def _coerce_team(value: Any) -> int | None:
    """Treat blank cells / placeholders as missing team numbers."""
    if value in (None, "", " "):
        return None
    try:
        n = int(value)
        if n <= 0:
            return None
        return n
    except (ValueError, TypeError):
        return None


def _classify_sheet(title: str | None, sheet_name: str) -> str:
    """Decide if a worksheet holds qual or practice matches.

    Inspects (a) the title cell at A1 and (b) the worksheet's tab name.
    Defaults to "qual" — historically the most common case and the
    safest fallback.
    """
    haystack = ((title or "") + " " + (sheet_name or "")).lower()
    for needle, kind in SHEET_KIND_HEURISTICS:
        if needle in haystack:
            return kind
    return "qual"


def parse_xlsx(content: bytes) -> dict[str, Any]:
    """Parse an XLSX schedule export into a dict shaped like the
    PDF-import pipeline output.

    Returns:
        {
            "matches":      [ {match_num, time, red[], blue[], red_surrogate[], blue_surrogate[]} ... ],
            "practice":     [ ... same shape ... ],   # may be empty
            "format_detected": "FMS xlsx (N qual, M practice)",
            "notes":        "<warnings, or empty>",
            "page_count":   <number of sheets>,
        }

    Raises ValueError if the workbook isn't recognisable as an FMS
    export (no team-position columns where expected).
    """
    try:
        # data_only=True so formula cells return their cached values,
        # not the formula text. read_only=True streams; faster on big
        # workbooks but we don't really need it for typical event sizes.
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as e:
        raise ValueError(f"Could not open XLSX file: {e}")

    qual_matches: list[dict] = []
    practice_matches: list[dict] = []
    notes: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # Title is at row index 0, col 0 (A1).
        title = rows[0][0] if rows[0] else None
        kind = _classify_sheet(title, sheet_name)

        # Find the header row — looking for "Time" + "Description" + at
        # least one team-position column. Allows minor format drift
        # (extra title rows, varied capitalisation) without breaking.
        header_idx = -1
        for i, row in enumerate(rows[:15]):  # check first 15 rows
            if not row:
                continue
            cells = [str(c).strip().lower() if c is not None else "" for c in row]
            if "time" in cells and "description" in cells:
                # Expect 6 team-position columns somewhere after
                team_col_count = sum(
                    1 for c in cells
                    if c.startswith("blue ") or c.startswith("red ")
                )
                if team_col_count >= 6:
                    header_idx = i
                    break

        if header_idx < 0:
            notes.append(
                f"Sheet '{sheet_name}' didn't match the FMS layout "
                f"(no Time/Description/Blue 1-3/Red 1-3 header row). Skipped."
            )
            continue

        # Resolve column positions from the header row — robust to
        # column reordering, which our exports don't do but other
        # tools' exports might. Indices:
        header_cells = [
            (str(c).strip().lower() if c is not None else "")
            for c in rows[header_idx]
        ]

        def col_idx(name: str) -> int:
            try:
                return header_cells.index(name)
            except ValueError:
                return -1

        c_time  = col_idx("time")
        c_desc  = col_idx("description")
        c_blue  = [col_idx(f"blue {i}") for i in (1, 2, 3)]
        c_red   = [col_idx(f"red {i}")  for i in (1, 2, 3)]

        if c_time < 0 or c_desc < 0 or any(i < 0 for i in c_blue + c_red):
            notes.append(
                f"Sheet '{sheet_name}' is missing required columns. Skipped."
            )
            continue

        # Walk data rows.
        target = practice_matches if kind == "practice" else qual_matches
        for row in rows[header_idx + 1:]:
            if not row or all(c is None or c == "" for c in row):
                continue

            # Pull description and extract match number.
            desc_val = row[c_desc] if c_desc < len(row) else None
            if desc_val is None:
                continue
            m = _MATCH_NUM_RE.search(str(desc_val))
            if not m:
                continue
            match_num = int(m.group(1))

            time_val = row[c_time] if c_time < len(row) else None
            time_str = str(time_val).strip() if time_val is not None else ""

            blue = [_coerce_team(row[i] if i < len(row) else None) for i in c_blue]
            red  = [_coerce_team(row[i] if i < len(row) else None) for i in c_red]

            # Skip rows that don't have at least one team — guards
            # against stray comment / footer rows getting picked up.
            if not any(blue) and not any(red):
                continue

            # Match-list shape consistent with pdf_extract pipeline.
            # Surrogate flags default to False; FMS export doesn't
            # encode them.
            target.append({
                "match_num":        match_num,
                "time":             time_str,
                "red":              [t for t in red  if t is not None],
                "blue":             [t for t in blue if t is not None],
                "red_surrogate":    [False] * len([t for t in red  if t is not None]),
                "blue_surrogate":   [False] * len([t for t in blue if t is not None]),
            })

    n_qual     = len(qual_matches)
    n_practice = len(practice_matches)

    if n_qual == 0 and n_practice == 0:
        raise ValueError(
            "No matches found. The workbook didn't match the expected "
            "FMS schedule format (Time / Description / Blue 1-3 / "
            "Red 1-3 columns)."
        )

    return {
        "matches":         qual_matches,
        "practice":        practice_matches,
        "format_detected": f"FMS xlsx ({n_qual} qual, {n_practice} practice)",
        "notes":           "; ".join(notes),
        "page_count":      len(wb.sheetnames),
    }
